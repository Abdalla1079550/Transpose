from __future__ import annotations

import os
import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

# Ensure app imports with DEMO_MODE fallback by default for local smoke.
os.environ.setdefault("DEMO_MODE", "1")

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.main import app  # noqa: E402


def _assert_ok(response, expected_status: int = 200) -> dict:
    if response.status_code != expected_status:
        raise RuntimeError(f"Expected {expected_status}, got {response.status_code}: {response.text}")
    return response.json()


def _build_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Mustafa Ahmed")
    document.add_paragraph("Data analyst candidate with Python, SQL, Tableau, and experimentation skills.")
    document.add_paragraph("Built demand forecasting dashboards and automated KPI reporting pipelines.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def main() -> None:
    with TestClient(app) as client:
        health = _assert_ok(client.get("/health"))
        snapshot = _assert_ok(
            client.get(
                "/market/snapshot",
                params={"region": "US", "role_cluster": "data_analyst", "days": 7},
            )
        )
        questions = _assert_ok(client.get("/students/interview-questions"))

        cv_bytes = _build_docx_bytes()
        create_student_resp = _assert_ok(
            client.post(
                "/students",
                data={
                    "full_name": "Mustafa Ahmed",
                    "email": "mustafa@example.com",
                    "location": "US",
                    "grad_year": "2026",
                    "skills": "python,sql,tableau,communication",
                },
                files={
                    "cv": (
                        "mustafa_cv.docx",
                        cv_bytes,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            ),
            expected_status=200,
        )
        student_id = create_student_resp["id"]

        interview = _assert_ok(
            client.post(
                f"/students/{student_id}/interview",
                json={
                    "answers": [
                        "I improved campaign ROI reporting by 18 percent using SQL automation.",
                        "I validate metrics by reconciling source systems and anomaly checks.",
                    ],
                    "notes": "Strong communicator; comfortable with stakeholder demos.",
                },
            )
        )
        card = _assert_ok(client.get(f"/students/{student_id}/card"))

        role = _assert_ok(
            client.post(
                "/roles",
                json={
                    "title": "Junior Data Analyst",
                    "description": "Support BI reporting, SQL analysis, and dashboard delivery.",
                    "location": "US",
                    "must_have_skills": ["python", "sql", "tableau"],
                    "min_grad_year": 2025,
                },
            )
        )
        role_id = role["id"]
        matches = _assert_ok(client.get(f"/roles/{role_id}/matches", params={"include_failed": "true"}))

    print("Smoke test passed")
    print(
        {
            "health": health,
            "market_source": snapshot.get("source"),
            "questions_count": len(questions.get("questions", [])),
            "student_id": student_id,
            "interview_saved": bool(interview.get("interview_notes")),
            "card_keys": sorted(card.get("card", {}).keys()),
            "role_id": role_id,
            "matches_total": matches.get("total"),
        }
    )


if __name__ == "__main__":
    main()
